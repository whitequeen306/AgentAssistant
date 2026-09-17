"""Deterministic Markdown → DOCX conversion for notes and research reports.

走程序、不走模型：报告本身是结构化 markdown，程序转换零 token、内容零漂移、
格式可复现；让模型"生成文档"又慢又贵，还会顺手润色掉辛苦核对的数据。

设计取舍：
- 只覆盖调研/笔记实际用到的 markdown 子集（标题 / 段落 / 表格 / 列表 /
  引用 / 代码块 / 分隔线 / 行内粗体·斜体·代码·链接），不追求完整规范；
- 超宽表（≥7 列，如择校对比表）自动切横排页——A4 竖排放 10 列必然挤成竖字；
- 链接转成真超链接（run 级），不再整段贴裸 URL。
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

_FRONTMATTER_RE = re.compile(r"\A---\s*\n.*?\n---\s*\n?", re.DOTALL)
# 行内标记：粗体 / 行内代码 / 链接 / 斜体（顺序即优先级，斜体放最后防吞 **）
_INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)\s]+\)|\*[^*\s][^*]*\*)")
_CELL_SPLIT_RE = re.compile(r"(?<!\\)\|")
_TABLE_ROW_RE = re.compile(r"^\s*\|.*\|\s*$")
_HR_RE = re.compile(r"^\s*(?:-{3,}|\*{3,}|_{3,})\s*$")
_SEP_CELL_RE = re.compile(r"^:?-{3,}:?$")
_IMG_RE = re.compile(r"^!\[([^\]]*)\]\(([^)\s]+)\)\s*$")

_BODY_FONT = "微软雅黑"
_MONO_FONT = "Consolas"
_TABLE_FONT_PT = 9
#: 列数达到该值即切横排页——再窄汉字就要竖着挤了
_LANDSCAPE_COLS = 7


def _clean_cell(text: str) -> str:
    """单元格内清理：还原转义管道、去掉行内标记（docx 表格里保留纯文本）。"""
    text = text.replace("\\|", "|").strip()
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)\s]+\)", r"\1", text)
    text = re.sub(r"\*([^*\s][^*]*)\*", r"\1", text)
    return text


def _parse_table_row(line: str) -> list[str]:
    parts = _CELL_SPLIT_RE.split(line.strip().strip("|"))
    return [_clean_cell(p) for p in parts]


def _is_separator_row(line: str) -> bool:
    cells = _parse_table_row(line)
    return bool(cells) and all(_SEP_CELL_RE.match(c) for c in cells if c != "")


def _set_east_asian(style_or_run: Any, font: str) -> None:
    """python-docx 只设 font.name 不影响中文，必须补 eastAsia。"""
    from docx.oxml.ns import qn

    rpr = style_or_run.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        from docx.oxml import OxmlElement

        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font)


def _add_hyperlink(paragraph: Any, url: str, text: str) -> None:
    """run 级真超链接（python-docx 没有高级 API，走 relationship）。"""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    part = paragraph.part
    r_id = part.relate_to(
        url, RT.HYPERLINK, is_external=True
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rPr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rPr.append(underline)
    run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _add_inline(paragraph: Any, text: str) -> None:
    """按行内标记拆分并写入 runs；链接转真超链接。"""
    for token in _INLINE_RE.split(text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = _MONO_FONT
            _set_east_asian(run, _MONO_FONT)
        elif token.startswith("["):
            m = re.match(r"\[([^\]]+)\]\(([^)\s]+)\)", token)
            if m:
                _add_hyperlink(paragraph, m.group(2), m.group(1))
            else:
                paragraph.add_run(token)
        elif token.startswith("*"):
            paragraph.add_run(token[1:]).italic = True
        else:
            paragraph.add_run(token)


def _strip_inline(text: str) -> str:
    """标题等场景只要纯文本。"""
    text = re.sub(r"\[([^\]]+)\]\([^)\s]+\)", r"\1", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\*([^*\s][^*]*)\*", r"\1", text)
    return text


def _shade_cell(cell: Any, hex_fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_fill)
    cell._tc.get_or_add_tcPr().append(shd)


def _switch_orientation(doc: Any, landscape: bool) -> None:
    from docx.enum.section import WD_ORIENT, WD_SECTION
    from docx.shared import Cm

    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    w, h = sec.page_width, sec.page_height
    if landscape and w < h:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = h, w
    elif not landscape and w > h:
        sec.orientation = WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = h, w
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    sec.left_margin = sec.right_margin = Cm(1.8)


def _add_table(doc: Any, rows: list[list[str]]) -> None:
    from docx.shared import Pt

    rows = [r for r in rows if any(c for c in r)]
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    wide = ncols >= _LANDSCAPE_COLS
    if wide:
        _switch_orientation(doc, landscape=True)

    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    table.autofit = True
    for i, row in enumerate(rows):
        for j in range(ncols):
            cell = table.cell(i, j)
            text = row[j] if j < len(row) else ""
            para = cell.paragraphs[0]
            para.paragraph_format.space_before = None
            para.paragraph_format.space_after = None
            run = para.add_run(text)
            run.font.size = Pt(8 if wide else _TABLE_FONT_PT)
            if i == 0:
                run.bold = True
                _shade_cell(cell, "EDEDED")
    # 表格后补一个空段，避免连续表格粘连
    doc.add_paragraph()


def _add_image(doc: Any, alt: str, src: str, base_dir: Path) -> None:
    """插入截图 + 居中题注；文件缺失时放显式占位而不是静默吞掉。"""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    path = Path(src)
    if not path.is_absolute():
        path = base_dir / path
    if path.exists():
        doc.add_picture(str(path), width=Cm(15))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        placeholder = doc.add_paragraph()
        placeholder.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = placeholder.add_run(f"[图片缺失：{src}]")
        run.font.color.rgb = None  # 保持默认色，仅文字提示
        run.italic = True
    if alt:
        caption = doc.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = caption.add_run(alt)
        run.font.size = Pt(9)
        run.italic = True


def markdown_to_docx(
    markdown_text: str,
    out_path: str | Path,
    *,
    title: str | None = None,
    base_dir: str | Path | None = None,
) -> Path:
    """Convert markdown to a formatted .docx. Returns the output path.

    ``base_dir`` is the resolution root for relative image paths;
    defaults to the output file's parent directory.
    """
    from docx import Document
    from docx.shared import Cm, Pt

    text = _FRONTMATTER_RE.sub("", markdown_text or "").replace("\r\n", "\n")
    root = Path(base_dir) if base_dir else Path(out_path).parent

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = _BODY_FONT
    normal.font.size = Pt(10.5)
    _set_east_asian(normal, _BODY_FONT)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.3

    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Cm(2.2)
        sec.left_margin = sec.right_margin = Cm(2.4)

    if title:
        doc.core_properties.title = title

    lines = text.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]

        if not line.strip():
            i += 1
            continue

        # 图片（截图进文档的关键；必须在段落合并之前处理）
        img = _IMG_RE.match(line)
        if img:
            _add_image(doc, img.group(1), img.group(2), root)
            i += 1
            continue

        # 代码块
        if line.lstrip().startswith("```"):
            block: list[str] = []
            i += 1
            while i < len(lines) and not lines[i].lstrip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1  # 跳过闭合围栏
            for code_line in block:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                run = p.add_run(code_line)
                run.font.name = _MONO_FONT
                run.font.size = Pt(9)
                _set_east_asian(run, _MONO_FONT)
            continue

        # 表格（含表头分隔行才算，避免把普通竖线段落误判成表）
        if _TABLE_ROW_RE.match(line) and i + 1 < len(lines) and _is_separator_row(lines[i + 1]):
            rows: list[list[str]] = [_parse_table_row(line)]
            i += 2
            while i < len(lines) and _TABLE_ROW_RE.match(lines[i]):
                rows.append(_parse_table_row(lines[i]))
                i += 1
            _add_table(doc, rows)
            continue

        # 分隔线（非表格行）
        if _HR_RE.match(line):
            i += 1
            continue

        # 标题
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            level = min(len(m.group(1)), 4)
            heading = doc.add_heading("", level=level)
            _add_inline(heading, _strip_inline(m.group(2)))
            for run in heading.runs:
                _set_east_asian(run, _BODY_FONT)
            i += 1
            continue

        # 引用
        if line.lstrip().startswith(">"):
            quote_lines: list[str] = []
            while i < len(lines) and lines[i].lstrip().startswith(">"):
                quote_lines.append(lines[i].lstrip().lstrip(">").strip())
                i += 1
            para = doc.add_paragraph()
            try:
                para.style = doc.styles["Intense Quote"]
            except KeyError:
                para.paragraph_format.left_indent = Cm(0.5)
            _add_inline(para, " ".join(q for q in quote_lines if q))
            continue

        # 列表（支持两级缩进）
        lm = re.match(r"^(\s*)([-*]|\d+[.)])\s+(.*)$", line)
        if lm:
            indent, marker, content = lm.group(1), lm.group(2), lm.group(3)
            ordered = marker[0].isdigit()
            style_name = "List Number" if ordered else "List Bullet"
            if len(indent) >= 2:
                style_name += " 2"
            try:
                para = doc.add_paragraph(style=style_name)
            except KeyError:
                para = doc.add_paragraph()
                para.paragraph_format.left_indent = Cm(0.75)
            _add_inline(para, content)
            i += 1
            continue

        # 普通段落（软换行合并到同一段，避免每行一个段导致间距稀碎）
        para_lines = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(
            r"^(#{1,6}\s|\s*[-*]\s|\s*\d+[.)]\s|>|\||```|!\[|\s*-{3,}\s*$)", lines[i]
        ):
            para_lines.append(lines[i])
            i += 1
        para = doc.add_paragraph()
        _add_inline(para, " ".join(s.strip() for s in para_lines))

    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out
