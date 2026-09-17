"""重建项目计划书 docx：md -> docx -> Token 黑色后处理 -> 覆盖桌面。"""
import copy
import pathlib
import shutil
import sys

ROOT = pathlib.Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from agent_assistant.export.md_to_docx import markdown_to_docx  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

SRC = ROOT / "competition" / "04-business-plan.md"
OUT = ROOT / "competition" / "项目计划书-词元新学（AgentAssistant）.docx"
DESKTOP = pathlib.Path("C:/Users/hp/Desktop/项目计划书-词元新学.docx")


def make_run(base_r, text, black=False):
    r = copy.deepcopy(base_r._element)
    for t in r.findall(qn("w:t")):
        r.remove(t)
    t = r.makeelement(qn("w:t"), {})
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    if black:
        rPr = r.find(qn("w:rPr"))
        if rPr is None:
            rPr = r.makeelement(qn("w:rPr"), {})
            r.insert(0, rPr)
        for c in rPr.findall(qn("w:color")):
            rPr.remove(c)
        c = rPr.makeelement(qn("w:color"), {})
        c.set(qn("w:val"), "000000")
        rPr.append(c)
    return r


def blacken_token(path):
    from docx import Document

    doc = Document(str(path))
    n = 0
    for p in doc.paragraphs[:12]:
        style = (p.style.name or "") if p.style is not None else ""
        text = p.text or ""
        if "Token" not in text:
            continue
        if "Heading" not in style and "Quote" not in style:
            continue
        runs = p.runs
        if not runs:
            continue
        base = runs[0]
        new_runs = []
        rest = text
        while "Token" in rest:
            i = rest.index("Token")
            if i:
                new_runs.append(make_run(base, rest[:i]))
            new_runs.append(make_run(base, "Token", black=True))
            rest = rest[i + 5:]
        if rest:
            new_runs.append(make_run(base, rest))
        for r in runs:
            r._element.getparent().remove(r._element)
        for r in new_runs:
            p._p.append(r)
        n += 1
    doc.save(str(path))
    return n


def insert_toc(path, levels="2-3"):
    """在第一处 Heading 2（正文第一章）之前插入原生 TOC 域目录。

    章标题是 Heading 2、节标题是 Heading 3，故默认 \o "2-3"，
    避免把封面两个 Heading 1 收进目录。
    同时置 w:updateFields=true，Word/WPS 打开时会提示更新并生成页码。
    """
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.oxml import OxmlElement
    from docx.shared import Pt

    doc = Document(str(path))
    target = None
    for p in doc.paragraphs:
        if p.style.name == "Heading 2":
            target = p
            break
    if target is None:
        return False

    # 占位提示（更新域后被真实目录替换）
    toc_p = doc.add_paragraph()
    r1 = toc_p.add_run()
    f_begin = OxmlElement("w:fldChar")
    f_begin.set(qn("w:fldCharType"), "begin")
    f_begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = rf'TOC \o "{levels}" \h \z \u'
    f_sep = OxmlElement("w:fldChar")
    f_sep.set(qn("w:fldCharType"), "separate")
    r1._r.append(f_begin)
    r1._r.append(instr)
    r1._r.append(f_sep)
    toc_p.add_run("（首次打开若目录为空：全选后按 F9，或右键“更新域”生成页码）")
    r3 = toc_p.add_run()
    f_end = OxmlElement("w:fldChar")
    f_end.set(qn("w:fldCharType"), "end")
    r3._r.append(f_end)

    # 目录标题（普通段落，避免被 TOC 收录）
    title_p = doc.add_paragraph()
    tr = title_p.add_run("目  录")
    tr.bold = True
    tr.font.size = Pt(16)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 目录后分页，正文从新页开始
    br_p = doc.add_paragraph()
    br_p.add_run().add_break(WD_BREAK.PAGE)

    # addprevious 每次紧邻插到 target 之前，故调用顺序即最终排列顺序
    target._p.addprevious(title_p._p)
    target._p.addprevious(toc_p._p)
    target._p.addprevious(br_p._p)

    uf = OxmlElement("w:updateFields")
    uf.set(qn("w:val"), "true")
    doc.settings.element.append(uf)
    doc.save(str(path))
    return True


def main():
    md = SRC.read_text(encoding="utf-8")
    markdown_to_docx(md, str(OUT), title="项目计划书", base_dir=str(ROOT / "competition"))
    n = blacken_token(OUT)
    insert_toc(OUT)
    shutil.copyfile(OUT, DESKTOP)
    print(f"token-black paragraphs: {n}")
    print(f"out: {OUT}")
    print(f"desktop: {DESKTOP}  size={DESKTOP.stat().st_size}")


if __name__ == "__main__":
    main()
